#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export Amazon Bulksheet cho "Lên camp" từ search term đã chọn.
Input (stdin JSON): { campaigns: LaunchCampaign[] } — xem lib/ppc/campaign-launch.ts
- adProduct "SP": sheet Sponsored Products Campaigns (MANUAL keyword / product targeting, hoặc AUTO)
- adProduct "SD": sheet Sponsored Display Campaigns (tactic T00020, contextual targeting theo ASIN)
Dùng template chuẩn: templates/ppc/AdvertisingBulksheetTemplate-seller.xlsx
"""

import sys
import json
import os
import io
import datetime
import warnings
warnings.filterwarnings("ignore")
import openpyxl

TEMPLATE_PATH = os.path.join(
    os.path.dirname(__file__), "..", "templates", "ppc", "AdvertisingBulksheetTemplate-seller.xlsx"
)
MATCH_TYPES = ("exact", "phrase", "broad")
BIDDING_STRATEGIES = ("Dynamic bids - down only", "Dynamic bids - up and down", "Fixed bid")


def money(value, fallback):
    try:
        v = float(value)
        return round(v, 2) if v > 0 else fallback
    except (TypeError, ValueError):
        return fallback


class SheetWriter:
    def __init__(self, ws, product):
        self.ws = ws
        self.product = product
        self.headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}
        self.row = ws.max_row + 1

    def write(self, values):
        for h, v in {"Product": self.product, **values}.items():
            col = self.headers.get(h)
            if col and v is not None and str(v).strip() != "":
                self.ws.cell(row=self.row, column=col, value=v)
        self.row += 1


def write_sp_campaign(w, camp, today_str):
    name = camp["name"]
    strategy = str(camp.get("biddingStrategy") or "").strip()
    w.write({
        "Entity": "Campaign",
        "Operation": "Create",
        "Campaign ID": name,
        "Campaign Name": name,
        "Start Date": today_str,
        "Targeting Type": "AUTO" if camp.get("targetingType") == "AUTO" else "MANUAL",
        "State": "enabled",
        "Daily Budget": money(camp.get("dailyBudget"), 5.0),
        "Bidding Strategy": strategy if strategy in BIDDING_STRATEGIES else BIDDING_STRATEGIES[0],
    })
    for group in camp.get("adGroups", []):
        ag = str(group.get("name") or name).strip()
        default_bid = money(group.get("defaultBid"), 1.0)
        ids = {"Campaign ID": name, "Ad Group ID": ag, "Campaign Name": name, "Ad Group Name": ag}
        w.write({"Entity": "Ad Group", "Operation": "Create", **ids, "State": "enabled", "Ad Group Default Bid": default_bid})
        for sku in group.get("skus", []):
            if str(sku or "").strip():
                w.write({"Entity": "Product Ad", "Operation": "Create", **ids, "SKU": str(sku).strip(), "State": "enabled"})
        for kw in group.get("keywords", []):
            text = str(kw.get("text") or "").strip()
            match = str(kw.get("matchType") or "exact").strip().lower()
            if text:
                w.write({
                    "Entity": "Keyword", "Operation": "Create", **ids, "State": "enabled",
                    "Bid": money(kw.get("bid"), default_bid), "Keyword Text": text,
                    "Match Type": match if match in MATCH_TYPES else "exact",
                })
        for target in group.get("productTargets", []):
            expr = str(target.get("expression") or "").strip()
            if expr:
                w.write({
                    "Entity": "Product Targeting", "Operation": "Create", **ids, "State": "enabled",
                    "Bid": money(target.get("bid"), default_bid), "Product Targeting Expression": expr,
                })


def write_sd_campaign(w, camp, today_str):
    name = camp["name"]
    w.write({
        "Entity": "Campaign",
        "Operation": "Create",
        "Campaign ID": name,
        "Campaign Name": name,
        "Start Date": today_str,
        "State": "enabled",
        "Tactic": "T00020",
        "Budget Type": "daily",
        "Budget": money(camp.get("dailyBudget"), 5.0),
        "Cost Type": "cpc",
    })
    for group in camp.get("adGroups", []):
        ag = str(group.get("name") or name).strip()
        default_bid = money(group.get("defaultBid"), 1.0)
        ids = {"Campaign ID": name, "Ad Group ID": ag, "Campaign Name": name, "Ad Group Name": ag}
        w.write({
            "Entity": "Ad Group", "Operation": "Create", **ids, "State": "enabled",
            "Ad Group Default Bid": default_bid, "Bid Optimization": "clicks",
        })
        for sku in group.get("skus", []):
            if str(sku or "").strip():
                w.write({"Entity": "Product Ad", "Operation": "Create", **ids, "SKU": str(sku).strip(), "State": "enabled"})
        for target in group.get("productTargets", []):
            expr = str(target.get("expression") or "").strip()
            if expr:
                w.write({
                    "Entity": "Contextual Targeting", "Operation": "Create", **ids, "State": "enabled",
                    "Bid": money(target.get("bid"), default_bid), "Targeting Expression": expr,
                })


def export_campaign_launch_bulksheet(data, output_path=None):
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Official Amazon template not found at {TEMPLATE_PATH}")

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    writers = {
        "SP": SheetWriter(wb["Sponsored Products Campaigns"], "Sponsored Products"),
        "SD": SheetWriter(wb["Sponsored Display Campaigns"], "Sponsored Display"),
    }
    today_str = datetime.datetime.now().strftime("%Y%m%d")

    for camp in data.get("campaigns", []):
        if not str(camp.get("name") or "").strip():
            continue
        camp["name"] = str(camp["name"]).strip()
        if camp.get("adProduct") == "SD":
            write_sd_campaign(writers["SD"], camp, today_str)
        else:
            write_sp_campaign(writers["SP"], camp, today_str)

    if output_path:
        wb.save(output_path)
    else:
        buf = io.BytesIO()
        wb.save(buf)
        sys.stdout.buffer.write(buf.getvalue())


if __name__ == "__main__":
    raw = sys.stdin.read()
    export_campaign_launch_bulksheet(json.loads(raw) if raw else {"campaigns": []})
