#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export Amazon Bulksheet using the official Amazon blank template:
templates/ppc/AdvertisingBulksheetTemplate-seller.xlsx
"""

import sys
import json
import os
import warnings
warnings.filterwarnings("ignore")
import openpyxl

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "templates", "ppc", "AdvertisingBulksheetTemplate-seller.xlsx")

def detect_ad_type(rec):
    for field in ("campaign_type", "campaignType", "adType", "ad_type"):
        val = str(rec.get(field) or "").strip().upper()
        if val in ("SP", "SB", "SD"):
            return val
    camp_name = str(rec.get("campaignName") or rec.get("campaign_name") or "").upper()
    if any(k in camp_name for k in [" SB", "SB0", "VIDEO", "BRANDS", "HSA"]):
        return "SB"
    if any(k in camp_name for k in [" SD", "SD0", "DISPLAY", "TACTIC"]):
        return "SD"
    return "SP"

def export_bulksheet(recommendations, output_path=None):
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Official Amazon template not found at {TEMPLATE_PATH}")

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    
    # Pre-map available worksheets, headers, and row pointers
    sheet_configs = {
        "SP": "Sponsored Products Campaigns",
        "SB": "Sponsored Brands Campaigns",
        "SD": "Sponsored Display Campaigns",
    }
    worksheets = {}
    headers_map = {}
    row_pointers = {}

    for ad_type_code, sheet_name in sheet_configs.items():
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            worksheets[ad_type_code] = ws
            headers_map[ad_type_code] = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}
            row_pointers[ad_type_code] = ws.max_row + 1

    for rec in recommendations:
        ad_type = detect_ad_type(rec)
        ws = worksheets.get(ad_type, worksheets.get("SP"))
        headers = headers_map.get(ad_type, headers_map.get("SP"))
        if not ws or not headers:
            continue

        current_row = row_pointers[ad_type]

        rec_type = str(rec.get("recType") or rec.get("actionType") or rec.get("action_type") or "").upper()
        target_type = str(rec.get("targetType") or rec.get("entityType") or rec.get("entity_type") or "KEYWORD").upper()
        
        rec_bid = None
        for k in ("final_value", "finalValue", "recommendedBid", "userFinalBid", "currentBid", "old_value"):
            if rec.get(k) is not None and str(rec.get(k)).strip() != "":
                rec_bid = rec.get(k)
                break

        keyword_text = str(rec.get("keyword") or rec.get("targetKeyword") or rec.get("target_keyword") or "")
        is_product_target = (
            target_type in ("PRODUCT", "PRODUCT TARGETING", "TARGETING") or
            "asin=" in keyword_text.lower() or
            "category=" in keyword_text.lower()
        )

        source_match_type = str(rec.get("matchType") or rec.get("match_type") or "").strip().lower()
        if source_match_type in ("exact", "phrase", "broad"):
            match_type = source_match_type
        else:
            match_type = "exact"

        operation = "Update"
        state = "enabled"
        bid_val = ""
        daily_budget = ""

        if rec_type in ("NEGATIVE_KEYWORD",):
            entity = "Negative Product Targeting" if is_product_target else "Negative Keyword"
            operation = "Create"
            match_type = "negativeExact"
            state = "enabled"
            bid_val = ""
        elif rec_type in ("UPDATE_BUDGET", "BUDGET_UPDATE"):
            entity = "Campaign"
            operation = "Update"
            state = "enabled"
            daily_budget = round(float(rec_bid), 2) if rec_bid is not None else ""
        elif rec_type in ("PAUSE_TARGET",):
            if ad_type == "SD":
                entity = "Contextual Targeting"
            else:
                entity = "Product Targeting" if is_product_target else "Keyword"
            operation = "Update"
            state = "paused"
            bid_val = ""
        elif rec_type in ("HARVEST_KEYWORD",):
            entity = "Product Targeting" if is_product_target else "Keyword"
            operation = "Create"
            match_type = "exact"
            state = "enabled"
            bid_val = round(float(rec_bid), 2) if rec_bid is not None else 1.00
        else:
            if ad_type == "SD":
                entity = "Contextual Targeting"
            else:
                entity = "Product Targeting" if is_product_target else "Keyword"
            operation = "Update"
            state = "enabled"
            bid_val = round(float(rec_bid), 2) if rec_bid is not None and str(rec_bid).strip() != "" else ""

        campaign_id = str(rec.get("campaignId") or rec.get("campaign_id") or "")
        ad_group_id = str(rec.get("adGroupId") or rec.get("ad_group_id") or "")
        target_id = str(rec.get("keywordId") or rec.get("targetId") or rec.get("target_id") or "")
        campaign_name = str(rec.get("campaignName") or rec.get("campaign_name") or "")
        ad_group_name = str(rec.get("adGroupName") or rec.get("ad_group_name") or "")

        if ad_type == "SB":
            values = {
                "Product": "Sponsored Brands",
                "Entity": entity,
                "Operation": operation,
                "Campaign ID": campaign_id,
                "Ad Group ID": ad_group_id if entity != "Campaign" else "",
                "Keyword ID": target_id if not is_product_target and entity == "Keyword" and operation == "Update" else "",
                "Product Targeting ID": target_id if is_product_target and operation == "Update" else "",
                "Campaign Name": campaign_name,
                "State": state,
                "Budget": daily_budget if entity == "Campaign" else "",
                "Bid": bid_val if entity in ("Keyword", "Product Targeting") else "",
                "Keyword Text": "" if is_product_target or entity != "Keyword" else keyword_text,
                "Match Type": "" if is_product_target or entity != "Keyword" else match_type,
                "Product Targeting Expression": keyword_text if is_product_target else "",
            }
        elif ad_type == "SD":
            values = {
                "Product": "Sponsored Display",
                "Entity": entity,
                "Operation": operation,
                "Campaign ID": campaign_id,
                "Ad Group ID": ad_group_id if entity != "Campaign" else "",
                "Targeting ID": target_id if entity != "Campaign" and operation == "Update" else "",
                "Campaign Name": campaign_name,
                "Ad Group Name": ad_group_name if entity != "Campaign" else "",
                "State": state,
                "Budget": daily_budget if entity == "Campaign" else "",
                "Bid": bid_val if entity != "Campaign" else "",
                "Targeting Expression": keyword_text if entity != "Campaign" else "",
            }
        else:
            # Default to Sponsored Products
            values = {
                "Product": "Sponsored Products",
                "Entity": entity,
                "Operation": operation,
                "Campaign ID": campaign_id,
                "Ad Group ID": ad_group_id if entity != "Campaign" else "",
                "Keyword ID": target_id if not is_product_target and entity == "Keyword" and operation == "Update" else "",
                "Product Targeting ID": target_id if is_product_target and operation == "Update" else "",
                "Campaign Name": campaign_name,
                "Ad Group Name": ad_group_name if entity != "Campaign" else "",
                "State": state,
                "Daily Budget": daily_budget if entity == "Campaign" else "",
                "Bid": bid_val if entity in ("Keyword", "Product Targeting") else "",
                "Keyword Text": "" if is_product_target or entity != "Keyword" else keyword_text,
                "Match Type": "" if is_product_target or entity != "Keyword" else match_type,
                "Product Targeting Expression": keyword_text if is_product_target else "",
            }

        for header, value in values.items():
            column = headers.get(header)
            if column and value is not None and str(value).strip() != "":
                ws.cell(row=current_row, column=column, value=value)
        row_pointers[ad_type] += 1

    if output_path:
        wb.save(output_path)
    else:
        # Save to stdout as binary
        import io
        buf = io.BytesIO()
        wb.save(buf)
        sys.stdout.buffer.write(buf.getvalue())

if __name__ == "__main__":
    if len(sys.argv) > 2:
        input_json_path = sys.argv[1]
        output_xlsx_path = sys.argv[2]
        with open(input_json_path, "r", encoding="utf-8") as f:
            recs = json.load(f)
        export_bulksheet(recs, output_xlsx_path)
    else:
        # Read from stdin
        raw = sys.stdin.read()
        recs = json.loads(raw) if raw else []
        export_bulksheet(recs)
