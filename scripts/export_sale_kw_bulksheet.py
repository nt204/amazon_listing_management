#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Export Amazon Bulksheet for "Lên Camp Sale KW" (Sale KW Campaign Launch)
using the official Amazon blank template: templates/ppc/AdvertisingBulksheetTemplate-seller.xlsx
"""

import sys
import json
import os
import datetime
import warnings
warnings.filterwarnings("ignore")
import openpyxl

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "templates", "ppc", "AdvertisingBulksheetTemplate-seller.xlsx")

def export_sale_kw_bulksheet(data, output_path=None):
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Official Amazon template not found at {TEMPLATE_PATH}")

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    sheet_name = "Sponsored Products Campaigns"
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in template.")

    ws = wb[sheet_name]
    headers = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}
    
    current_row = ws.max_row + 1
    today_str = datetime.datetime.now().strftime("%Y%m%d")

    campaigns = data.get("campaigns", []) if isinstance(data, dict) else data

    for camp in campaigns:
        target_camp_name = str(camp.get("targetCampaignName") or "").strip()
        source_camp_name = str(camp.get("sourceCampaignName") or "").strip()
        if not target_camp_name:
            continue

        ad_group_name = str(camp.get("adGroupName") or target_camp_name).strip()
        daily_budget = float(camp.get("dailyBudget") or 10.0)
        default_bid = float(camp.get("defaultBid") or 1.0)
        bidding_strategy = str(camp.get("biddingStrategy") or "Dynamic bids - down only").strip()
        skus_raw = camp.get("sku") or camp.get("skus") or []
        if isinstance(skus_raw, str):
            skus = [s.strip() for s in skus_raw.split(",") if s.strip()]
        else:
            skus = [str(s).strip() for s in skus_raw if str(s).strip()]

        source_camp_id = str(camp.get("sourceCampaignId") or "").strip()
        source_ag_id = str(camp.get("sourceAdGroupId") or "").strip()
        source_ag_name = str(camp.get("sourceAdGroupName") or source_camp_name).strip()
        negate_in_source = bool(camp.get("negateInSource", False))

        keywords = camp.get("keywords", [])

        # 1. Row: Campaign Create
        campaign_row = {
            "Product": "Sponsored Products",
            "Entity": "Campaign",
            "Operation": "Create",
            "Campaign ID": target_camp_name,
            "Campaign Name": target_camp_name,
            "Start Date": today_str,
            "Targeting Type": "MANUAL",
            "State": "enabled",
            "Daily Budget": round(daily_budget, 2),
            "Bidding Strategy": bidding_strategy,
        }
        for h, v in campaign_row.items():
            col = headers.get(h)
            if col and v is not None and str(v).strip() != "":
                ws.cell(row=current_row, column=col, value=v)
        current_row += 1

        # 2. Row: Ad Group Create
        ad_group_row = {
            "Product": "Sponsored Products",
            "Entity": "Ad Group",
            "Operation": "Create",
            "Campaign ID": target_camp_name,
            "Ad Group ID": ad_group_name,
            "Campaign Name": target_camp_name,
            "Ad Group Name": ad_group_name,
            "State": "enabled",
            "Ad Group Default Bid": round(default_bid, 2),
        }
        for h, v in ad_group_row.items():
            col = headers.get(h)
            if col and v is not None and str(v).strip() != "":
                ws.cell(row=current_row, column=col, value=v)
        current_row += 1

        # 3. Row(s): Product Ad Create
        if skus:
            for s in skus:
                ad_row = {
                    "Product": "Sponsored Products",
                    "Entity": "Product Ad",
                    "Operation": "Create",
                    "Campaign ID": target_camp_name,
                    "Ad Group ID": ad_group_name,
                    "Campaign Name": target_camp_name,
                    "Ad Group Name": ad_group_name,
                    "SKU": s,
                    "State": "enabled",
                }
                for h, v in ad_row.items():
                    col = headers.get(h)
                    if col and v is not None and str(v).strip() != "":
                        ws.cell(row=current_row, column=col, value=v)
                current_row += 1

        # 4. Row(s): Keywords & Targets Create
        for kw in keywords:
            term = str(kw.get("customerSearchTerm") or kw.get("keyword") or "").strip()
            if not term:
                continue

            kw_bid = kw.get("bid")
            try:
                bid_val = round(float(kw_bid), 2) if kw_bid is not None else round(default_bid, 2)
            except (ValueError, TypeError):
                bid_val = round(default_bid, 2)

            is_product = (
                term.lower().startswith("b0") or
                term.lower().startswith("asin=") or
                term.lower().startswith("category=")
            )

            if is_product:
                asin_clean = term.strip()
                expr = asin_clean if asin_clean.lower().startswith("asin=") else f'asin="{asin_clean.upper()}"'
                target_row = {
                    "Product": "Sponsored Products",
                    "Entity": "Product Targeting",
                    "Operation": "Create",
                    "Campaign ID": target_camp_name,
                    "Ad Group ID": ad_group_name,
                    "Campaign Name": target_camp_name,
                    "Ad Group Name": ad_group_name,
                    "State": "enabled",
                    "Bid": bid_val,
                    "Product Targeting Expression": expr,
                }
            else:
                match_t = str(kw.get("matchType") or "exact").strip().lower()
                if match_t not in ("exact", "phrase", "broad"):
                    match_t = "exact"
                target_row = {
                    "Product": "Sponsored Products",
                    "Entity": "Keyword",
                    "Operation": "Create",
                    "Campaign ID": target_camp_name,
                    "Ad Group ID": ad_group_name,
                    "Campaign Name": target_camp_name,
                    "Ad Group Name": ad_group_name,
                    "State": "enabled",
                    "Bid": bid_val,
                    "Keyword Text": term,
                    "Match Type": match_t,
                }

            for h, v in target_row.items():
                col = headers.get(h)
                if col and v is not None and str(v).strip() != "":
                    ws.cell(row=current_row, column=col, value=v)
            current_row += 1

        # 5. (Optional) Phủ định Negative Exact ở Campaign gốc
        if negate_in_source and source_camp_id:
            for kw in keywords:
                term = str(kw.get("customerSearchTerm") or kw.get("keyword") or "").strip()
                if not term:
                    continue

                is_product = (
                    term.lower().startswith("b0") or
                    term.lower().startswith("asin=") or
                    term.lower().startswith("category=")
                )

                if is_product:
                    asin_clean = term.strip()
                    expr = asin_clean if asin_clean.lower().startswith("asin=") else f'asin="{asin_clean.upper()}"'
                    neg_row = {
                        "Product": "Sponsored Products",
                        "Entity": "Negative Product Targeting" if source_ag_id else "Campaign Negative Product Targeting",
                        "Operation": "Create",
                        "Campaign ID": source_camp_id,
                        "Ad Group ID": source_ag_id if source_ag_id else "",
                        "Campaign Name": source_camp_name,
                        "Ad Group Name": source_ag_name if source_ag_id else "",
                        "State": "enabled",
                        "Product Targeting Expression": expr,
                    }
                else:
                    neg_row = {
                        "Product": "Sponsored Products",
                        "Entity": "Negative Keyword" if source_ag_id else "Campaign Negative Keyword",
                        "Operation": "Create",
                        "Campaign ID": source_camp_id,
                        "Ad Group ID": source_ag_id if source_ag_id else "",
                        "Campaign Name": source_camp_name,
                        "Ad Group Name": source_ag_name if source_ag_id else "",
                        "State": "enabled",
                        "Keyword Text": term,
                        "Match Type": "negativeExact",
                    }

                for h, v in neg_row.items():
                    col = headers.get(h)
                    if col and v is not None and str(v).strip() != "":
                        ws.cell(row=current_row, column=col, value=v)
                current_row += 1

    if output_path:
        wb.save(output_path)
    else:
        import io
        buf = io.BytesIO()
        wb.save(buf)
        sys.stdout.buffer.write(buf.getvalue())

if __name__ == "__main__":
    if len(sys.argv) > 2:
        input_json_path = sys.argv[1]
        output_xlsx_path = sys.argv[2]
        with open(input_json_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        export_sale_kw_bulksheet(data, output_xlsx_path)
    else:
        raw = sys.stdin.read()
        data = json.loads(raw) if raw else {"campaigns": []}
        export_sale_kw_bulksheet(data)
