import warnings
import pandas as pd
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
import re
import logging
import time
import requests
from urllib.parse import quote
from multiprocessing import Pool, freeze_support
import pytz
import gspread
from oauth2client.service_account import ServiceAccountCredentials

# --- CẤU HÌNH ---

# Bỏ qua các cảnh báo không cần thiết từ openpyxl
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Thiết lập logging
VN_TZ = pytz.timezone('Asia/Ho_Chi_Minh')
current_date = datetime.now(VN_TZ).strftime("%Y%m%d")
log_file = f"E:\\PythonProject1\\Amzdata-04-bulkfile-{current_date}.txt"

# Xóa log cũ trước khi chạy
if os.path.exists(log_file):
    with open(log_file, 'w', encoding='utf-8') as f:
        f.write('')

logging.basicConfig(
    filename=log_file,
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S',
    encoding='utf-8'
)
console = logging.StreamHandler()
console.setLevel(logging.INFO)
console.setFormatter(logging.Formatter('%(asctime)s - %(levelname)s - %(message)s', datefmt='%Y-%m-%d %H:%M:%S'))
logging.getLogger('').addHandler(console)

# Cấu hình Telegram
BOT_TOKEN = "7319056820:AAGnHjcTzqWlS6QU4OjEiP3mPnXCejGnXqw"
GROUP_CHAT_ID = -1002462247213
TOPIC_ID = 1151
TELEGRAM_PROXY_URL = "https://nce-telegram-proxy.ntthanhnce.workers.dev"
SCRIPT_NAME = "Amz-04-bulkfile"

# Định nghĩa đường dẫn
# input_base_directory = r"\\Nceglobal\design tong hop\1. DỮ LIỆU PPC\DỮ LIỆU ĐẦU VÀO"
input_base_directory = r"E:\PPC\DỮ LIỆU ĐẦU VÀO"
# output_base_directory = r"\\Nceglobal\design tong hop\1. DỮ LIỆU PPC\DỮ LIỆU ĐẦU RA"
output_base_directory = r"E:\PPC\DỮ LIỆU ĐẦU RA"

# Cấu hình Google Sheet API
GOOGLE_API_KEYFILE = r"E:\PythonProject1\NCE_googleapi.json"
SPREADSHEET_URL = "https://docs.google.com/spreadsheets/d/1CBZuKZOjG9zxNrfdh3qgeofkwW1758ruZl-yT77_IeI/edit?gid=0#gid=0"

# Cấu hình signature
USE_SIGNATURE = False
SIGNATURE = "Thanh"

# Lấy danh sách store FBM tự động từ Google Sheet
try:
    scope = [
        "https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive",
        "https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive.file",
    ]
    creds = ServiceAccountCredentials.from_json_keyfile_name(GOOGLE_API_KEYFILE, scope)
    client = gspread.authorize(creds)
    sheet = client.open_by_url(SPREADSHEET_URL).worksheet("APBrand")
    all_records = sheet.get_all_records()
    fbm_stores = [
        row["BRAND"] for row in all_records
        if str(row.get("FBM", "")).strip().upper() == "YES" and row.get("BRAND")
    ]
    logging.info(f"[SYSTEM] Đã lấy danh sách FBM tự động từ Sheet: {fbm_stores}")
except Exception as e:
    logging.error(f"[SYSTEM] Lỗi khi lấy danh sách FBM từ Sheet: {e}. Sẽ chạy với danh sách rỗng.")
    fbm_stores = []


# --- CÁC HÀM HỖ TRỢ ---

def sanitize_filename(name):
    """Loại bỏ các ký tự không hợp lệ cho tên file/thư mục."""
    return re.sub(r'[\\/*?:"<>|]', ' ', str(name))


def escape_markdown(text):
    """Hàm hỗ trợ escape các ký tự đặc biệt của MarkdownV2."""
    escape_chars = r'_*[]()~`>#+-=|{}.!'
    return ''.join(f'\\{char}' if char in escape_chars else char for char in str(text))


def clean_for_log(markdown_text):
    """Làm sạch tin nhắn Markdown để ghi log cho dễ đọc."""
    emojis_to_remove = ['🚀', '✅', '❌', '🔑', '💥', '🎉', '⚠️', '☠️', '📄', '⏳']
    clean_text = markdown_text
    for emoji in emojis_to_remove:
        clean_text = clean_text.replace(emoji, '')
    chars_to_remove = ['*', '_', '`', '~']
    for char in chars_to_remove:
        clean_text = clean_text.replace(char, '')
    clean_text = clean_text.replace('\\', '')
    return clean_text.strip()


def send_telegram_message(message, parse_mode="MarkdownV2"):
    """Gửi tin nhắn Telegram. Tự retry khi gặp 429 (rate limit)."""
    url = f"{TELEGRAM_PROXY_URL}/bot{BOT_TOKEN}/sendMessage?chat_id={GROUP_CHAT_ID}&text={quote(message)}&message_thread_id={TOPIC_ID}&parse_mode={parse_mode}"
    for attempt in range(1, 4):
        try:
            response = requests.get(url, timeout=10)
            if response.status_code == 200:
                clean_message = clean_for_log(message)
                log_message = clean_message.replace('\n', ' ').strip()
                logging.info(f"[TELEGRAM] Đã gửi: {log_message}")
                return
            if response.status_code == 429:
                try:
                    retry_after = int(response.json().get("parameters", {}).get("retry_after", 5))
                except Exception:
                    retry_after = 5
                if attempt < 3:
                    logging.warning(
                        f"[TELEGRAM] 429 rate limit, chờ {retry_after}s rồi thử lại (lần {attempt}/3).")
                    time.sleep(retry_after + 1)
                    continue
            logging.error(f"[TELEGRAM] Gửi thất bại: {response.status_code} {response.text}")
            return
        except Exception as e:
            logging.error(f"[TELEGRAM] Lỗi khi gửi tin nhắn: {str(e)}")
            return


def format_duration(seconds):
    minutes, seconds = divmod(seconds, 60)
    return f"{int(minutes)} phút {int(seconds)} giây"


def get_on_stores(keyfile, spreadsheet_url):
    """Lấy danh sách các store có Status là 'ON' từ Google Sheet."""
    try:
        scope = [
            "https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive",
            "https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive.file",
        ]
        creds = ServiceAccountCredentials.from_json_keyfile_name(keyfile, scope)
        client = gspread.authorize(creds)
        sheet = client.open_by_url(spreadsheet_url).worksheet("APBrand")
        records = sheet.get_all_records()
        on_stores = {row['BRAND'] for row in records if str(row.get('Status', '')).strip().upper() == 'ON'}
        logging.info(f"[SYSTEM] Các store có status 'ON' từ Google Sheet: {on_stores}")
        return on_stores
    except Exception as e:
        logging.error(f"[SYSTEM] Lỗi khi lấy danh sách store 'ON': {e}. Sẽ xử lý tất cả store tìm thấy.")
        return None


# --- CÁC HÀM CHỨC NĂNG ---

def write_excel_with_retry(df, output_file, campaign_type, retries=3, delay=10):
    for attempt in range(retries):
        try:
            output_dir = os.path.dirname(output_file)
            os.makedirs(output_dir, exist_ok=True)
            df.to_excel(output_file, index=False, engine='openpyxl')
            apply_formatting(output_file, campaign_type)
            return True
        except PermissionError:
            logging.warning(
                f"Permission denied khi ghi {output_file}. Thử lại lần {attempt + 1}/{retries} sau {delay} giây...")
            msg = f"*{escape_markdown(SCRIPT_NAME)}* \\- _Lỗi ghi file_\n" \
                  f"📄 _File_: `{escape_markdown(os.path.basename(output_file))}`\n" \
                  f"⚠️ _Lỗi Permission Denied, đang thử lại\\.\\.\\._"
            send_telegram_message(msg)
            time.sleep(delay)
            if attempt == retries - 1:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                new_output_file = f"{os.path.splitext(output_file)[0]}_{timestamp}.xlsx"
                logging.error(f"Không thể ghi file {output_file} sau {retries} lần thử. Lưu vào {new_output_file}")
                msg_fail = f"*{escape_markdown(SCRIPT_NAME)}* \\- _Lỗi ghi file_\n" \
                           f"📄 _File_: `{escape_markdown(os.path.basename(output_file))}`\n" \
                           f"❌ _Thất bại sau {retries} lần thử. Đã lưu vào file mới: `{escape_markdown(os.path.basename(new_output_file))}`_"
                send_telegram_message(msg_fail)
                try:
                    df.to_excel(new_output_file, index=False, engine='openpyxl')
                    apply_formatting(new_output_file, campaign_type)
                    return True
                except Exception as e2:
                    logging.error(f"Lỗi khi ghi file thay thế {new_output_file}: {str(e2)}")
                    msg_final_fail = f"*{escape_markdown(SCRIPT_NAME)}* \\- _Lỗi ghi file_\n" \
                                     f"💥 _Ghi file thay thế cũng thất bại: `{escape_markdown(os.path.basename(new_output_file))}`_"
                    send_telegram_message(msg_final_fail)
                    return False
    return False


def create_output_directory(output_base, store_name, range_day, fulfillment, report_type):
    folder_name = f"{current_date}-{SIGNATURE}" if USE_SIGNATURE else current_date
    date_directory = os.path.join(output_base, folder_name)
    os.makedirs(date_directory, exist_ok=True)
    store_directory = os.path.join(date_directory, store_name)
    os.makedirs(store_directory, exist_ok=True)
    range_directory = os.path.join(store_directory, range_day)
    os.makedirs(range_directory, exist_ok=True)
    current_time = datetime.now().strftime("%Y-%m-%d")
    final_directory = os.path.join(range_directory, f"[{fulfillment}] Bulk File {report_type} {current_time}")
    os.makedirs(final_directory, exist_ok=True)
    return final_directory


def filter_and_save_excel(input_file, output_directory, range_day_report, store_name, fulfillment, report_type):
    df_combined = pd.DataFrame()
    try:
        xls = pd.ExcelFile(input_file, engine='openpyxl')
    except Exception as e:
        logging.error(f"[{store_name}] Lỗi khi mở file {os.path.basename(input_file)}: {e}")
        msg = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_name)}*\n" \
              f"❌ _Lỗi mở file {fulfillment} {report_type}_ `{escape_markdown(os.path.basename(input_file))}`\n" \
              f"Lỗi: `{escape_markdown(e)}`"
        send_telegram_message(msg)
        return

    # --- START: SỬA LỖI HOA/THƯỜNG ---
    sheet_name_keyword = "sponsored brands campaigns" if report_type == "SB" else "sponsored products campaigns"
    sheet_names = [sheet for sheet in xls.sheet_names if sheet_name_keyword in sheet.lower()]
    # --- END: SỬA LỖI HOA/THƯỜNG ---

    if not sheet_names:
        logging.info(
            f"[{store_name}] Không tìm thấy sheet chứa '{sheet_name_keyword}' trong file: {os.path.basename(input_file)}")
        return

    for sheet_name in sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet_name)

        # --- START: SỬA LỖI HOA/THƯỜNG ---
        # Chuẩn hóa tên cột về chữ thường
        df.columns = [str(col).strip().lower() for col in df.columns]

        # Các cột cần thiết (đã được chuyển thành chữ thường)
        required_cols = ['portfolio name (informational only)', 'campaign name (informational only)']
        if not all(col in df.columns for col in required_cols):
            logging.warning(f"[{store_name}] Bỏ qua sheet '{sheet_name}' vì thiếu cột portfolio hoặc campaign name")
            continue
        # --- END: SỬA LỖI HOA/THƯỜNG ---

        if 'operation' in df.columns:
            df['operation'] = 'Update'
        df_combined = pd.concat([df_combined, df], ignore_index=True)

    if df_combined.empty or 'portfolio name (informational only)' not in df_combined.columns:
        logging.info(
            f"[{store_name}] Không có dữ liệu {report_type} hợp lệ hoặc thiếu cột Portfolio Name trong file: {os.path.basename(input_file)}")
        return

    df_combined['portfolio name (informational only)'] = df_combined['portfolio name (informational only)'].fillna(
        '').astype(str)
    skus = sorted(df_combined['portfolio name (informational only)'].unique())
    if not skus:
        output_file = os.path.join(output_directory, f"no_sku_data_{range_day_report}.xlsx")
        write_excel_with_retry(df_combined, output_file, None)
        return

    for sku in skus:
        safe_sku = sanitize_filename(sku)
        is_fbm_sku = "FBM" in sku.upper()
        if (fulfillment == "FBM" and not is_fbm_sku) or (fulfillment == "FBA" and is_fbm_sku):
            continue
        if "OFF" in sku.upper():
            continue

        sku_data = df_combined[df_combined['portfolio name (informational only)'] == sku]

        if report_type == "SB":
            collection_dir = os.path.join(output_directory, "SB01")
            os.makedirs(collection_dir, exist_ok=True)
            video_dir = os.path.join(output_directory, "SB05")
            os.makedirs(video_dir, exist_ok=True)

            specific_match_types = ['Exact', 'Phrase', 'Broad']
            filter_condition = sku_data['match type'].isin(specific_match_types)
            collection_condition = sku_data['campaign name (informational only)'].str.contains(
                r'\b(?:sb01|Collection)\b', case=False, regex=True, na=False)
            collection_data = sku_data[collection_condition & filter_condition]
            if not collection_data.empty:
                output_file = os.path.join(collection_dir, f"{safe_sku} Bulk Collection {range_day_report}.xlsx")
                write_excel_with_retry(collection_data.sort_values(by='spend', ascending=False), output_file,
                                       'Collection')

            video_condition = sku_data['campaign name (informational only)'].str.contains(r'\b(?:sb05|video)\b',
                                                                                          case=False, regex=True,
                                                                                          na=False)
            video_data = sku_data[video_condition & filter_condition]
            if not video_data.empty:
                output_file = os.path.join(video_dir, f"{safe_sku} Bulk Video {range_day_report}.xlsx")
                write_excel_with_retry(video_data.sort_values(by='spend', ascending=False), output_file, 'Video')

        elif report_type == "SP":
            if fulfillment == "FBM":
                auto_dir = os.path.join(output_directory, "AUTO")
                keyword_dir = os.path.join(output_directory, "FBM")
                os.makedirs(auto_dir, exist_ok=True)
                os.makedirs(keyword_dir, exist_ok=True)
            else:
                auto_dir = os.path.join(output_directory, "SP04 (Auto)")
                keyword_dir = os.path.join(output_directory, "SP03")
                os.makedirs(auto_dir, exist_ok=True)
                os.makedirs(keyword_dir, exist_ok=True)

            filter_condition_kw = sku_data['campaign name (informational only)'].str.contains('Exact|Phrase|Broad',
                                                                                              case=False, na=False)
            filter_condition_kw &= sku_data['match type'].isin(['Exact', 'Phrase', 'Broad'])
            filtered_data_kw = sku_data[filter_condition_kw]
            if not filtered_data_kw.empty:
                output_file = os.path.join(keyword_dir, f"{safe_sku} Bulk File SP {range_day_report}.xlsx")
                write_excel_with_retry(filtered_data_kw.sort_values(by='spend', ascending=False), output_file,
                                       'Keyword')

            if 'product targeting expression' in sku_data.columns:
                filter_condition_auto = sku_data['product targeting expression'].str.contains(
                    'substitutes|loose-match|complements|close-match', case=False, na=False)
                filter_condition_auto &= sku_data['product targeting expression'].isin(
                    ['substitutes', 'loose-match', 'complements', 'close-match'])
                filtered_data_auto = sku_data[filter_condition_auto]
                if not filtered_data_auto.empty:
                    output_file = os.path.join(auto_dir, f"{safe_sku} Bulk File SP {range_day_report} Auto.xlsx")
                    write_excel_with_retry(filtered_data_auto.sort_values(by='spend', ascending=False), output_file,
                                           'Auto')


def apply_formatting(file_path, campaign_type=None):
    try:
        workbook = load_workbook(file_path)
        sheet = workbook.active
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions

        green_fill = PatternFill(start_color='228b22', end_color='228b22', fill_type='solid')
        white_green_fill = PatternFill(start_color='8fbc8f', end_color='8fbc8f', fill_type='solid')
        white_red_fill = PatternFill(start_color='e9967a', end_color='e9967a', fill_type='solid')
        red_fill = PatternFill(start_color='ff0000', end_color='ff0000', fill_type='solid')
        brown_fill = PatternFill(start_color='cd853f', end_color='cd853f', fill_type='solid')

        header = [cell.value.lower() if cell.value else '' for cell in sheet[1]]
        orders_col_index = acos_col_index = clicks_col_index = None
        if 'orders' in header:
            orders_col_index = header.index('orders') + 1
        if 'acos' in header:
            acos_col_index = header.index('acos') + 1
        if 'clicks' in header:
            clicks_col_index = header.index('clicks') + 1

        if orders_col_index and acos_col_index and clicks_col_index:
            for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
                orders_cell = row[orders_col_index - 1]
                acos_cell = row[acos_col_index - 1]
                clicks_cell = row[clicks_col_index - 1]
                if orders_cell.value is not None and orders_cell.value > 0:
                    if acos_cell.value is not None:
                        fill = None
                        if 0 < acos_cell.value <= 0.2:
                            fill = green_fill
                        elif 0.2 < acos_cell.value <= 0.4:
                            fill = white_green_fill
                        elif 0.4 < acos_cell.value <= 0.6:
                            fill = white_red_fill
                        else:
                            fill = red_fill
                        if fill:
                            for cell in row: cell.fill = fill
                else:
                    click_threshold = 9 if campaign_type == 'Collection' else 8 if campaign_type == 'Video' else 7
                    if clicks_cell.value is not None and clicks_cell.value > click_threshold:
                        for cell in row: cell.fill = brown_fill
        workbook.save(file_path)
    except Exception as e:
        logging.error(f"Lỗi khi định dạng file {file_path}: {str(e)}")
        msg = f"*{escape_markdown(SCRIPT_NAME)}* \\- _Lỗi định dạng file_\n" \
              f"📄 _File_: `{escape_markdown(os.path.basename(file_path))}`\n" \
              f"Lỗi: `{escape_markdown(e)}`"
        send_telegram_message(msg)


def process_store(args):
    store_folder, date_input_directory, output_base = args
    store_path = os.path.join(date_input_directory, store_folder)
    if not os.path.isdir(store_path): return

    logging.info(f"--- Bắt đầu xử lý Bulkfile cho store: {store_folder} ---")
    msg = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
          f"🚀 _Bắt đầu xử lý Bulkfile_"
    send_telegram_message(msg)
    store_start_time = time.time()

    try:
        if store_folder in fbm_stores:
            sp_fbm_start_time = time.time()
            msg_sp_fbm_start = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
                               f"⏳ _Bắt đầu xử lý Bulk File SP FBM\\.\\.\\._"
            send_telegram_message(msg_sp_fbm_start)
            for file in [f for f in os.listdir(store_path) if "Bulk File SP" in f]:
                input_file_path = os.path.join(store_path, file)
                parts = file.split(" ")
                range_day_report = parts[-2][1:] + " " + parts[-1][:-6]
                output_dir_fbm_sp = create_output_directory(output_base, store_folder, range_day_report, "FBM", "SP")
                filter_and_save_excel(input_file_path, output_dir_fbm_sp, range_day_report, store_folder, "FBM", "SP")
            sp_fbm_end_time = time.time()
            msg_sp_fbm_end = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
                             f"✅ _Hoàn thành xử lý Bulk File SP FBM_\n" \
                             f"Thời gian: `{format_duration(sp_fbm_end_time - sp_fbm_start_time)}`"
            send_telegram_message(msg_sp_fbm_end)

            sb_fbm_start_time = time.time()
            msg_sb_fbm_start = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
                               f"⏳ _Bắt đầu xử lý Bulk File SB FBM\\.\\.\\._"
            send_telegram_message(msg_sb_fbm_start)
            for file in [f for f in os.listdir(store_path) if "Bulk File SB" in f]:
                input_file_path = os.path.join(store_path, file)
                parts = file.split(" ")
                range_day_report = parts[-2][1:] + " " + parts[-1][:-6]
                output_dir_fbm_sb = create_output_directory(output_base, store_folder, range_day_report, "FBM", "SB")
                filter_and_save_excel(input_file_path, output_dir_fbm_sb, range_day_report, store_folder, "FBM", "SB")
            sb_fbm_end_time = time.time()
            msg_sb_fbm_end = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
                             f"✅ _Hoàn thành xử lý Bulk File SB FBM_\n" \
                             f"Thời gian: `{format_duration(sb_fbm_end_time - sb_fbm_start_time)}`"
            send_telegram_message(msg_sb_fbm_end)

        sp_fba_start_time = time.time()
        msg_sp_fba_start = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
                           f"⏳ _Bắt đầu xử lý Bulk File SP FBA\\.\\.\\._"
        send_telegram_message(msg_sp_fba_start)
        for file in [f for f in os.listdir(store_path) if "Bulk File SP" in f]:
            input_file_path = os.path.join(store_path, file)
            parts = file.split(" ")
            range_day_report = parts[-2][1:] + " " + parts[-1][:-6]
            output_dir_fba_sp = create_output_directory(output_base, store_folder, range_day_report, "FBA", "SP")
            filter_and_save_excel(input_file_path, output_dir_fba_sp, range_day_report, store_folder, "FBA", "SP")
        sp_fba_end_time = time.time()
        msg_sp_fba_end = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
                         f"✅ _Hoàn thành xử lý Bulk File SP FBA_\n" \
                         f"Thời gian: `{format_duration(sp_fba_end_time - sp_fba_start_time)}`"
        send_telegram_message(msg_sp_fba_end)

        sb_fba_start_time = time.time()
        msg_sb_fba_start = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
                           f"⏳ _Bắt đầu xử lý Bulk File SB FBA\\.\\.\\._"
        send_telegram_message(msg_sb_fba_start)
        for file in [f for f in os.listdir(store_path) if "Bulk File SB" in f]:
            input_file_path = os.path.join(store_path, file)
            parts = file.split(" ")
            range_day_report = parts[-2][1:] + " " + parts[-1][:-6]
            output_dir_fba_sb = create_output_directory(output_base, store_folder, range_day_report, "FBA", "SB")
            filter_and_save_excel(input_file_path, output_dir_fba_sb, range_day_report, store_folder, "FBA", "SB")
        sb_fba_end_time = time.time()
        msg_sb_fba_end = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
                         f"✅ _Hoàn thành xử lý Bulk File SB FBA_\n" \
                         f"Thời gian: `{format_duration(sb_fba_end_time - sb_fba_start_time)}`"
        send_telegram_message(msg_sb_fba_end)

        store_end_time = time.time()
        store_duration = store_end_time - store_start_time
        logging.info(
            f"--- Hoàn thành xử lý Bulkfile cho store {store_folder} trong {format_duration(store_duration)} ---")
        msg = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
              f"🎉 _Hoàn thành xử lý Bulkfile_\n" \
              f"Tổng thời gian: `{format_duration(store_duration)}`"
        send_telegram_message(msg)
    except Exception as e:
        logging.error(f"[{store_folder}] Lỗi nghiêm trọng khi xử lý store: {e}")
        msg = f"*{escape_markdown(SCRIPT_NAME)}* \\- *{escape_markdown(store_folder)}*\n" \
              f"💥 _Lỗi nghiêm trọng khi xử lý store_\n" \
              f"Lỗi: `{escape_markdown(e)}`"
        send_telegram_message(msg)


if __name__ == '__main__':
    freeze_support()
    start_msg = f"🚀 *{escape_markdown(SCRIPT_NAME)}* \\- _Bắt đầu chạy script_"
    send_telegram_message(start_msg)
    logging.info(f"=============================================")
    logging.info(f"BẮT ĐẦU CHẠY SCRIPT: {SCRIPT_NAME}")
    logging.info(f"=============================================")

    overall_start_time = time.time()

    date_input_directory = os.path.join(input_base_directory, current_date)
    if not os.path.exists(date_input_directory):
        logging.warning(f"Thư mục đầu vào cho ngày hôm nay không tồn tại: {date_input_directory}")
        msg = f"⚠️ *{escape_markdown(SCRIPT_NAME)}* \\- _Không tìm thấy thư mục dữ liệu cho ngày hôm nay. Dừng script._"
        send_telegram_message(msg)
        exit()

    on_stores_from_sheet = get_on_stores(GOOGLE_API_KEYFILE, SPREADSHEET_URL)

    all_available_folders = [folder for folder in os.listdir(date_input_directory) if
                             os.path.isdir(os.path.join(date_input_directory, folder))]

    if on_stores_from_sheet is not None:
        store_folders_to_process = [
            folder for folder in all_available_folders if folder in on_stores_from_sheet
        ]
        logging.info(f"Các store sẽ được xử lý (status ON và có thư mục): {store_folders_to_process}")
    else:
        store_folders_to_process = all_available_folders

    args_list = [(folder, date_input_directory, output_base_directory) for folder in store_folders_to_process]

    if not args_list:
        logging.warning("Không có store nào để xử lý sau khi lọc theo status 'ON'.")
    else:
        with Pool(processes=6) as pool:
            pool.map(process_store, args_list)

    overall_end_time = time.time()
    overall_duration = overall_end_time - overall_start_time
    logging.info(f"TỔNG THỜI GIAN XỬ LÝ TẤT CẢ STORES: {format_duration(overall_duration)}")
    end_msg = f"✅ *{escape_markdown(SCRIPT_NAME)}* \\- _Hoàn thành xử lý Bulkfile cho tất cả stores_\n" \
              f"Tổng thời gian: `{format_duration(overall_duration)}`"
    send_telegram_message(end_msg)
